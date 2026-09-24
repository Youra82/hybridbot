#!/usr/bin/env python3
"""
run_portfolio_optimizer.py  (hybridbot)

1:1 aus zerobot/run_portfolio_optimizer.py uebernommen: laedt alle
market_sense-Configs, fuehrt Portfolio-Simulation durch und waehlt das beste
Portfolio per Greedy-Algorithmus (portfolio_optimizer.py). Schreibt
active_strategies in settings.json.

Aufruf:
  python3 run_portfolio_optimizer.py              # interaktiv
  python3 run_portfolio_optimizer.py --auto-write  # automatisch (Scheduler)
"""
import os
import sys
import json
import argparse
import pandas as pd
from datetime import date, timedelta
from tqdm import tqdm

PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(PROJECT_ROOT, 'src'))

from hybridbot.analysis.backtester import fetch_ohlcv_history

CONFIGS_DIR = os.path.join(PROJECT_ROOT, 'src', 'hybridbot', 'strategy', 'configs')
SETTINGS_PATH = os.path.join(PROJECT_ROOT, 'settings.json')
OOS_FILE = os.path.join(PROJECT_ROOT, 'artifacts', 'results', 'last_oos_run.json')

B = '\033[1;37m'
G = '\033[0;32m'
Y = '\033[1;33m'
R = '\033[0;31m'
NC = '\033[0m'

DEFAULT_LOOKBACK_DAYS = 730
BOT_NAME = 'hybridbot'


def _scan_configs() -> list:
    if not os.path.isdir(CONFIGS_DIR):
        return []
    return sorted([os.path.join(CONFIGS_DIR, f) for f in os.listdir(CONFIGS_DIR) if f.endswith('.json')])


def _load_oos_info() -> tuple:
    """
    oos_start    = MAX(alle per-Config _meta.oos_start, last_oos_run.json oos_start)
                   -> konservativster gemeinsamer Startpunkt (kein Lookahead)
    warmup_start = MIN(alle per-Config _meta.train_start, last_oos_run.json warmup_start)
    oos_map      = pro Config: OOS-Ergebnis aus last_oos_run.json (fuer PnL-Filter)
    """
    oos_starts, warmup_starts, oos_map = [], [], {}

    if os.path.isdir(CONFIGS_DIR):
        for fn in os.listdir(CONFIGS_DIR):
            if not (fn.startswith('config_') and fn.endswith('.json')):
                continue
            try:
                with open(os.path.join(CONFIGS_DIR, fn)) as f:
                    cfg = json.load(f)
                meta = cfg.get('_meta', {})
                if meta.get('oos_start'):
                    oos_starts.append(meta['oos_start'])
                if meta.get('train_start'):
                    warmup_starts.append(meta['train_start'])
            except Exception:
                pass

    if os.path.exists(OOS_FILE):
        try:
            with open(OOS_FILE) as f:
                data = json.load(f)
            if data.get('oos_start'):
                oos_starts.append(data['oos_start'])
            if data.get('warmup_start'):
                warmup_starts.append(data['warmup_start'])
            for r in data.get('results', []):
                cfg_file = r.get('config_file', '')
                if cfg_file:
                    oos_map[cfg_file] = r
        except Exception:
            pass

    if not oos_starts:
        return None, None, oos_map

    oos_start = max(oos_starts)
    warmup_start = min(warmup_starts) if warmup_starts else None

    if len(set(oos_starts)) > 1:
        print(f"  {Y}Verschiedene OOS-Startdaten erkannt:{NC}")
        for d in sorted(set(oos_starts)):
            print(f"    {d}")
        print(f"  -> Gemeinsamer OOS-Start: {B}{oos_start}{NC} (spaetester = sicherster)")

    return oos_start, warmup_start, oos_map


def _build_strategies_data(config_files: list, fallback_start: str, end_date: str) -> dict:
    strategies_data = {}
    for path in tqdm(config_files, desc='Lade Configs & Daten'):
        fname = os.path.basename(path)
        try:
            with open(path) as f:
                config = json.load(f)
            market = config.get('market', {})
            symbol = market.get('symbol', '')
            timeframe = market.get('timeframe', '')
            if not symbol or not timeframe:
                continue
            # Per-Config Warmup aus _meta.train_start -- exakt wie oos_tester.py.
            per_start = config.get('_meta', {}).get('train_start', fallback_start)
            df = fetch_ohlcv_history(symbol, timeframe, per_start, end_date, 'bitget')
            if df is None or df.empty or len(df) < 50:
                print(f"  {Y}Uebersprungen (keine Daten): {fname}{NC}")
                continue

            risk = config.get('risk', {})
            strategies_data[fname] = {
                'symbol': symbol, 'timeframe': timeframe, 'df': df,
                'market_sense': config.get('market_sense', {}),
                'risk_per_trade_pct': risk.get('risk_per_trade_pct', 1.0),
                'leverage': risk.get('leverage', 10),
            }
        except Exception as e:
            print(f"  {Y}Fehler bei {fname}: {e}{NC}")
    return strategies_data


def _write_to_settings(portfolio_files: list, strategies_data: dict) -> None:
    with open(SETTINGS_PATH) as f:
        settings = json.load(f)
    existing = settings.get('live_trading_settings', {}).get('active_strategies', [])
    existing_map = {(s.get('symbol'), s.get('timeframe')): s for s in existing}
    new_strategies = []
    for fname in portfolio_files:
        sd = strategies_data.get(fname, {})
        symbol, timeframe = sd.get('symbol', ''), sd.get('timeframe', '')
        if not symbol or not timeframe:
            continue
        base = existing_map.get((symbol, timeframe), {})
        new_strategies.append({**base, 'symbol': symbol, 'timeframe': timeframe, 'active': True})
    lt = settings.setdefault('live_trading_settings', {})
    lt['active_strategies'] = new_strategies
    lt['use_auto_optimizer_results'] = True
    with open(SETTINGS_PATH, 'w') as f:
        json.dump(settings, f, indent=4)


def _get_telegram_creds():
    try:
        with open(os.path.join(PROJECT_ROOT, 'secret.json')) as f:
            s = json.load(f)
        tg = s.get('telegram', {})
        t, c = tg.get('bot_token', ''), tg.get('chat_id', '')
        return (t, c) if t and c else (None, None)
    except Exception:
        return None, None


def _send_telegram(msg):
    token, chat = _get_telegram_creds()
    if not token:
        return
    try:
        import requests
        requests.post(f'https://api.telegram.org/bot{token}/sendMessage',
                     data={'chat_id': chat, 'text': msg}, timeout=10)
    except Exception:
        pass


def _send_telegram_doc(fpath, caption=''):
    token, chat = _get_telegram_creds()
    if not token:
        return
    try:
        import requests
        with open(fpath, 'rb') as fh:
            requests.post(f'https://api.telegram.org/bot{token}/sendDocument',
                         data={'chat_id': chat, 'caption': caption},
                         files={'document': fh}, timeout=30)
    except Exception:
        pass


PAIR_COLORS = ['#f59e0b', '#8b5cf6', '#ec4899', '#14b8a6', '#f97316', '#84cc16', '#06b6d4', '#a78bfa']


def generate_equity_html(final, capital, data_start, end_date, labels,
                         portfolio_files=None, strategies_data=None):
    try:
        import plotly.graph_objects as go
        from plotly.subplots import make_subplots
    except ImportError:
        print(f'  {Y}plotly nicht installiert -- Chart uebersprungen.{NC}')
        return None

    eq_df = final.get('equity_curve')
    trades = final.get('trade_history', [])
    if eq_df is None or (hasattr(eq_df, 'empty') and eq_df.empty):
        return None

    eq_times = [str(t) for t in (eq_df['timestamp'] if 'timestamp' in eq_df.columns else eq_df.index)]
    eq_vals = [float(v) for v in eq_df['equity']]

    pnl = final.get('total_pnl_pct', 0)
    dd = final.get('max_drawdown_pct', 0)
    wr = final.get('win_rate', 0)
    n = final.get('trade_count', 0)
    eq = final.get('end_capital', eq_vals[-1] if eq_vals else capital)
    sign = '+' if pnl >= 0 else ''
    pairs_str = ', '.join(labels) if labels else BOT_NAME

    title = (f"{BOT_NAME} Portfolio -- {pairs_str} | Trades: {n} | WR: {wr:.1f}% | "
            f"PnL: {sign}{pnl:.1f}% | Equity: {eq:.2f} USDT | MaxDD: {dd:.1f}%")

    fig = make_subplots(specs=[[{"secondary_y": True}]])

    if trades and strategies_data:
        strategy_pairs = [(sd.get('symbol', ''), sd.get('timeframe', ''))
                          for sd in strategies_data.values() if sd.get('symbol') and sd.get('timeframe')]
        for idx, (sym, tf) in enumerate(strategy_pairs):
            pair_trades = sorted(
                [t for t in trades if t.get('symbol') == sym and t.get('timeframe') == tf],
                key=lambda t: t.get('entry_time', ''))
            if not pair_trades:
                continue
            peq = capital
            ptimes = [pair_trades[0].get('entry_time', '')[:19]]
            pvals = [peq]
            for t in pair_trades:
                peq += t.get('pnl', 0)
                ptimes.append(t.get('ts', t.get('entry_time', ''))[:19])
                pvals.append(round(peq, 2))
            fig.add_trace(go.Scatter(
                x=ptimes, y=pvals, mode='lines', name=f"{sym.split('/')[0]}/{tf}",
                line=dict(color=PAIR_COLORS[idx % len(PAIR_COLORS)], width=1), opacity=0.55,
            ), secondary_y=False)

    fig.add_hline(y=capital, line=dict(color='rgba(100,100,100,0.35)', width=1, dash='dash'),
                 annotation_text=f'Start {capital:.0f} USDT', annotation_position='top left')

    win_x, win_y, loss_x, loss_y = [], [], [], []
    if trades and eq_df is not None:
        eq_idx = eq_df.copy()
        eq_idx.index = pd.to_datetime(eq_idx.index, utc=True, errors='coerce')
        for t in trades:
            try:
                ts_dt = pd.to_datetime(t.get('ts', ''), utc=True, errors='coerce')
                if pd.isna(ts_dt):
                    continue
                idx_pos = eq_idx.index.get_indexer([ts_dt], method='nearest')[0]
                eq_val = float(eq_idx['equity'].iloc[idx_pos])
            except Exception:
                continue
            (win_x if t.get('pnl', 0) > 0 else loss_x).append(t.get('ts', '')[:19])
            (win_y if t.get('pnl', 0) > 0 else loss_y).append(eq_val)

    fig.add_trace(go.Scatter(x=eq_times, y=eq_vals, mode='lines', name='Portfolio Equity',
                             line=dict(color='#2563eb', width=2), opacity=0.75), secondary_y=True)
    if win_x:
        fig.add_trace(go.Scatter(x=win_x, y=win_y, mode='markers', name='Exit Gewinn',
                                 marker=dict(color='#22d3ee', symbol='circle', size=10,
                                            line=dict(width=1, color='#0e7490'))), secondary_y=True)
    if loss_x:
        fig.add_trace(go.Scatter(x=loss_x, y=loss_y, mode='markers', name='Exit Verlust',
                                 marker=dict(color='#ef4444', symbol='x', size=10,
                                            line=dict(width=2, color='#7f1d1d'))), secondary_y=True)

    fig.update_layout(
        title=dict(text=title, font=dict(size=13), x=0.5, xanchor='center'),
        height=750, hovermode='x unified', template='plotly_white', dragmode='zoom',
        xaxis=dict(rangeslider=dict(visible=True), fixedrange=False),
        legend=dict(orientation='h', yanchor='bottom', y=1.02, xanchor='center', x=0.5),
    )
    fig.update_yaxes(title_text='Einzel-Equity (USDT)', secondary_y=False, fixedrange=False)
    fig.update_yaxes(title_text='Portfolio-Equity (USDT)', secondary_y=True, fixedrange=False)

    from datetime import datetime as _dt
    outfile = f'/tmp/{BOT_NAME}_portfolio_{_dt.now().strftime("%Y%m%d_%H%M%S")}.html'
    fig.write_html(outfile)
    print(f'  {G}Chart erstellt: {outfile}{NC}')
    return outfile


def generate_trades_excel(final, capital, data_start, end_date, portfolio_files, strategies_data):
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print(f'  {Y}openpyxl nicht installiert -- Excel uebersprungen. (pip install openpyxl){NC}')
        return None

    trades = final.get('trade_history', [])
    if not trades:
        print(f'  {Y}Keine Trades -- Excel uebersprungen.{NC}')
        return None

    rows = []
    for i, t in enumerate(trades, 1):
        pnl = t.get('pnl', 0.0)
        sym = t.get('symbol', '')
        rows.append({
            'Nr': i,
            'Datum (Entry)': str(t.get('entry_time', ''))[:16].replace('T', ' '),
            'Datum (Exit)': str(t.get('ts', ''))[:16].replace('T', ' '),
            'Coin': sym.split('/')[0] if sym else '?',
            'Timeframe': t.get('timeframe', ''),
            'Richtung': t.get('direction', '?').upper(),
            'Entry': round(t.get('entry', 0.0), 6),
            'Exit': round(t.get('exit', 0.0), 6),
            'Ergebnis': 'Gewinn' if pnl > 0 else 'Verlust',
            'Margin (USDT)': round(t.get('margin_used', 0.0), 4),
            'PnL (USDT)': round(pnl, 4),
        })

    equity = capital
    for r in rows:
        equity += r['PnL (USDT)']
        r['Gesamtkapital'] = round(equity, 4)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Trades'

    header_fill = PatternFill('solid', fgColor='1E3A5F')
    win_fill = PatternFill('solid', fgColor='D6F4DC')
    loss_fill = PatternFill('solid', fgColor='FAD7D7')
    alt_fill = PatternFill('solid', fgColor='F2F2F2')
    thin_border = Border(*(Side(style='thin', color='CCCCCC') for _ in range(4)))
    col_widths = {'Nr': 6, 'Datum (Entry)': 18, 'Datum (Exit)': 18, 'Coin': 10, 'Timeframe': 12,
                 'Richtung': 10, 'Entry': 14, 'Exit': 14, 'Ergebnis': 12, 'Margin (USDT)': 16,
                 'PnL (USDT)': 14, 'Gesamtkapital': 16}

    headers = list(rows[0].keys())
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = col_widths.get(h, 14)
    ws.row_dimensions[1].height = 22

    for r_idx, row in enumerate(rows, 2):
        fill = win_fill if row['Ergebnis'] == 'Gewinn' else (loss_fill if r_idx % 2 == 0 else alt_fill)
        for col, key in enumerate(headers, 1):
            cell = ws.cell(row=r_idx, column=col, value=row[key])
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            if key in ('Entry', 'Exit', 'Margin (USDT)', 'PnL (USDT)', 'Gesamtkapital'):
                cell.number_format = '#,##0.0000'
        ws.row_dimensions[r_idx].height = 18

    pnl_total, wr, dd = final.get('total_pnl_pct', 0), final.get('win_rate', 0), final.get('max_drawdown_pct', 0)
    n, eq = final.get('trade_count', len(rows)), final.get('end_capital', equity)
    sign = '+' if pnl_total >= 0 else ''

    summary_row = len(rows) + 3
    ws.cell(row=summary_row, column=1, value='Zusammenfassung').font = Font(bold=True, size=11)
    for label, value in [('Trades gesamt', n), ('Win-Rate', f"{wr:.1f}%"), ('PnL', f"{sign}{pnl_total:.1f}%"),
                         ('Final Equity', f"{eq:.2f} USDT"), ('Max Drawdown', f"{dd:.1f}%"),
                         ('Zeitraum', f"{data_start} -> {end_date}")]:
        ws.cell(row=summary_row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=summary_row, column=2, value=value)
        summary_row += 1

    from datetime import datetime as _dt
    outfile = f'/tmp/{BOT_NAME}_trades_{_dt.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    wb.save(outfile)
    print(f'  {G}Excel-Tabelle erstellt: {outfile}{NC}')
    return outfile


def main() -> int:
    parser = argparse.ArgumentParser(description='hybridbot Portfolio-Optimizer')
    parser.add_argument('--capital', type=float, default=None)
    parser.add_argument('--max-dd', type=float, default=30.0)
    parser.add_argument('--start-date', type=str, default=None)
    parser.add_argument('--end-date', type=str, default=None)
    parser.add_argument('--auto-write', action='store_true')
    args = parser.parse_args()

    from hybridbot.analysis.portfolio_optimizer import run_portfolio_optimizer

    with open(SETTINGS_PATH) as f:
        settings = json.load(f)
    opt = settings.get('optimization_settings', {})
    capital = args.capital or float(opt.get('start_capital', 100))
    max_dd = args.max_dd
    end_date = args.end_date or date.today().strftime('%Y-%m-%d')
    max_positions = int(settings.get('live_trading_settings', {}).get('max_open_positions', 10))

    oos_start, warmup_start, oos_map = _load_oos_info()
    data_start = warmup_start or args.start_date or (date.today() - timedelta(days=DEFAULT_LOOKBACK_DAYS)).strftime('%Y-%m-%d')

    trade_start = oos_start
    lookback_weeks = opt.get('backtest_lookback_weeks')
    if lookback_weeks:
        lookback_start = (date.today() - timedelta(weeks=int(lookback_weeks))).strftime('%Y-%m-%d')
        trade_start = max(trade_start, lookback_start) if trade_start else lookback_start

    print(f"\n{'-'*72}")
    print(f"{B}  hybridbot -- Portfolio-Optimizer (market_sense){NC}")
    print(f"  Kapital: {capital:.0f} USDT | MaxDD <= {max_dd:.0f}%")
    if oos_start:
        print(f"  {Y}OOS-Periode:  {oos_start} -> {end_date}  (DUNKLER BEREICH){NC}")
        print(f"  Warmup:       {data_start} -> {oos_start}")
        if lookback_weeks and trade_start != oos_start:
            print(f"  {Y}Lookback:     letzte {lookback_weeks} Wochen ({trade_start} -> {end_date}){NC}")
        print(f"  OOS-Filter:   Nur Strategien mit positivem OOS-PnL werden zugelassen")
    else:
        print(f"  {Y}Kein last_oos_run.json gefunden -- nutze Zeitraum ohne OOS-Filter{NC}")
        print(f"  Zeitraum: {data_start} -> {end_date}")
    print(f"{'-'*72}\n")

    config_files = _scan_configs()
    if not config_files:
        print(f"{R}  Keine Configs in {CONFIGS_DIR}{NC}")
        print(f"  -> Zuerst den Optimizer ausfuehren!\n")
        return 1

    print(f"  {len(config_files)} Config(s) gefunden.\n")
    strategies_data = _build_strategies_data(config_files, data_start, end_date)
    if not strategies_data:
        print(f"{R}  Keine Daten geladen.{NC}")
        return 1

    smoothing_step_days = int(opt.get('smoothing_step_days', 2))
    smoothing_samples = int(opt.get('smoothing_samples', 7))

    result = run_portfolio_optimizer(
        capital, strategies_data, data_start, end_date, max_dd,
        trade_start_date=trade_start, oos_map=oos_map if oos_map else None,
        smoothing_step_days=smoothing_step_days, smoothing_samples=smoothing_samples,
        max_positions=max_positions)

    if not result or not result.get('optimal_portfolio'):
        print(f"{R}  Kein Portfolio erfuellt MaxDD <= {max_dd:.0f}%.{NC}\n")
        return 0

    # max_positions wird jetzt SCHON im Greedy-Loop durchgesetzt (siehe
    # portfolio_optimizer.py) -- optimal_portfolio ist daher nie laenger als
    # max_positions. Der Slice bleibt als reine Absicherung stehen, veraendert
    # aber nichts mehr an final_result (das war der eigentliche Bug: PnL/DD
    # kamen bisher vom UNGEKUERZTEN Portfolio, nicht vom hier angezeigten).
    portfolio_files = result['optimal_portfolio'][:max_positions]
    final = result.get('final_result') or {}

    print(f"\n{'='*72}")
    print(f"{B}  Optimales Portfolio -- {len(portfolio_files)} Strategie(n){NC}\n")
    for fname in portfolio_files:
        sd = strategies_data.get(fname, {})
        print(f"  {G}+{NC} {sd.get('symbol', fname):<26} / {sd.get('timeframe', ''):<6}")
    if final:
        print(f"\n  Endkapital: {final.get('end_capital', 0):.2f} USDT  "
             f"| PnL: {final.get('total_pnl_pct', 0):+.1f}%  | MaxDD: {final.get('max_drawdown_pct', 0):.2f}%")
    print(f"{'='*72}\n")

    labels = [f"{strategies_data.get(f, {}).get('symbol', '?')}/{strategies_data.get(f, {}).get('timeframe', '?')}"
             for f in portfolio_files]

    if args.auto_write:
        _write_to_settings(portfolio_files, strategies_data)
        print(f"{G}settings.json aktualisiert -- {len(portfolio_files)} Strategie(n).{NC}\n")
        if final:
            pnl, dd = final.get('total_pnl_pct', 0), final.get('max_drawdown_pct', 0)
            n, wr, eq = final.get('trade_count', 0), final.get('win_rate', 0), final.get('end_capital', 0)
            _send_telegram(f"{BOT_NAME} Auto-Optimizer\n{len(portfolio_files)} Strategien | {n} Trades | WR: {wr:.1f}%\n"
                          f"PnL: {pnl:+.1f}% | MaxDD: {dd:.1f}% | Equity: {eq:.2f} USDT\nZeitraum: {data_start} -> {end_date}")
            html = generate_equity_html(final, capital, data_start, end_date, labels, portfolio_files, strategies_data)
            if html:
                _send_telegram_doc(html, caption=f'{BOT_NAME} Portfolio-Equity | PnL: {pnl:+.1f}%')
            excel = generate_trades_excel(final, capital, data_start, end_date, portfolio_files, strategies_data)
            if excel:
                _send_telegram_doc(excel, caption=f'{BOT_NAME} Trades-Tabelle | {len(portfolio_files)} Strategien | '
                                                  f'{n} Trades | WR: {wr:.1f}% | Final: {eq:.2f} USDT')
    else:
        try:
            ans = input("  Portfolio in settings.json eintragen? (j/n): ").strip().lower()
        except (EOFError, KeyboardInterrupt):
            ans = 'n'
        if ans in ('j', 'ja', 'y', 'yes'):
            _write_to_settings(portfolio_files, strategies_data)
            print(f"{G}settings.json aktualisiert.{NC}\n")
        else:
            print(f"{Y}  settings.json NICHT geaendert.{NC}\n")

        try:
            chart_ans = input("  Interaktive Charts erstellen & via Telegram senden? (j/n): ").strip().lower()
        except (EOFError, KeyboardInterrupt):
            chart_ans = 'n'
        if chart_ans in ('j', 'ja', 'y', 'yes') and final:
            pnl, n, wr, eq = final.get('total_pnl_pct', 0), final.get('trade_count', 0), final.get('win_rate', 0), final.get('end_capital', 0)
            html = generate_equity_html(final, capital, data_start, end_date, labels, portfolio_files, strategies_data)
            if html:
                _send_telegram_doc(html, caption=f'{BOT_NAME} Portfolio-Equity | PnL: {pnl:+.1f}% | Equity: {eq:.2f} USDT | '
                                                 f'MaxDD: {final.get("max_drawdown_pct", 0):.1f}%')
            excel = generate_trades_excel(final, capital, data_start, end_date, portfolio_files, strategies_data)
            if excel:
                _send_telegram_doc(excel, caption=f'{BOT_NAME} Trades-Tabelle | {len(portfolio_files)} Strategien | '
                                                  f'{n} Trades | WR: {wr:.1f}% | Final: {eq:.2f} USDT')

    return 0


if __name__ == '__main__':
    sys.exit(main())
